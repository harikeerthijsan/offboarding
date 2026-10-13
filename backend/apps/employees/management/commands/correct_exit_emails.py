"""Correct exited employee/login emails from the Exit worksheet by JSAN ID."""
from collections import Counter

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook

from apps.employees.models import Employee


class Command(BaseCommand):
    help = 'Correct exited employee and login emails from an Excel Exit sheet.'

    def add_arguments(self, parser):
        parser.add_argument('workbook')
        parser.add_argument('--apply', action='store_true', help='Save changes (default is preview).')
        parser.add_argument('--skip-unresolved', action='store_true', help='Skip invalid/missing emails and ownership conflicts.')

    def handle(self, *args, **options):
        User = get_user_model()
        workbook = load_workbook(options['workbook'], read_only=True, data_only=True)
        try:
            sheet = workbook['Exit']
            headers = [str(v or '').strip() for v in next(sheet.iter_rows(values_only=True))]
            id_col = headers.index('JSANID')
            email_col = headers.index('Jsan Email ID')
            entries = [(str(r[id_col] or '').strip().upper(), str(r[email_col] or '').strip().rstrip('.').lower())
                       for r in sheet.iter_rows(min_row=2, values_only=True) if r[id_col]]
        finally:
            workbook.close()
        if any(n > 1 for n in Counter(eid for eid, _ in entries).values()):
            raise CommandError('Duplicate employee IDs in Exit sheet.')
        if any(n > 1 for n in Counter(email for _, email in entries).values()):
            raise CommandError('Duplicate emails in Exit sheet.')
        changed = 0
        with transaction.atomic():
            for eid, email in entries:
                try:
                    validate_email(email)
                    employee = Employee.objects.select_for_update().get(employee_id=eid)
                except (ValidationError, Employee.DoesNotExist) as exc:
                    if options['skip_unresolved']:
                        self.stdout.write(f'{eid}: skipped invalid email or missing employee')
                        continue
                    raise CommandError(f'{eid}: invalid email or missing employee; no changes saved.') from exc
                if employee.employment_status != 'EXITED':
                    raise CommandError(f'{eid}: not EXITED; no changes saved.')
                user = User.objects.select_for_update().get(pk=employee.user_id)
                if (Employee.objects.filter(email__iexact=email).exclude(pk=employee.pk).exists()
                        or User.objects.filter(email__iexact=email).exclude(pk=user.pk).exists()):
                    if options['skip_unresolved']:
                        self.stdout.write(f'{eid}: skipped email ownership conflict')
                        continue
                    raise CommandError(f'{eid}: email belongs to another record; no changes saved.')
                if employee.email == email and user.email == email:
                    continue
                self.stdout.write(f'{eid}: {employee.email} -> {email}')
                employee.email = email
                employee.save(update_fields=['email', 'updated_at'])
                # Login uses email; leave username and password intact.
                user.email = email
                user.save(update_fields=['email'])
                changed += 1
            self.stdout.write(f'{"Updated" if options["apply"] else "Would update"} {changed} of {len(entries)} exited employees.')
            if not options['apply']:
                transaction.set_rollback(True)
