"""Excel export of skill submissions — ported from the source app's lib/export.ts.

Two sheets (JSAN/QF/09 style): an Employee Skill Matrix (one row per employee,
latest matching submission) and Skill details (one row per skill). Built with
openpyxl. The embedded QF/09 logo image is omitted; the format box text is kept.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

FORMAT_BOX = "Format No: JSAN/QF/09\nRev : 1.0\nEffective date: 05-Apr-2024"
_thin = Side(style='thin')
_BORDER = Border(top=_thin, left=_thin, bottom=_thin, right=_thin)
_HDR_FONT = Font(name='Calibri', size=10, bold=True)
_BODY_FONT = Font(name='Calibri', size=10)
_WRAP = Alignment(vertical='center', wrap_text=True)
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _uniq(values):
    seen, out = set(), []
    for v in values:
        v = str(v or '').strip()
        if v and v.lower() not in seen:
            seen.add(v.lower())
            out.append(v)
    return out


def _total_years(skills):
    return sum(s.get('years') for s in skills if isinstance(s.get('years'), (int, float))) or None


def _tenure(doj, today=None):
    if not doj:
        return None, ''
    today = today or date.today()
    months = (today.year - doj.year) * 12 + (today.month - doj.month)
    if today.day < doj.day:
        months -= 1
    months = max(0, months)
    y, m = divmod(months, 12)
    return round((months / 12) * 10) / 10, f"{y} yr{'' if y == 1 else 's'} {m} mth{'' if m == 1 else 's'}"


def _skill_line(k):
    tail = ', '.join(x for x in [k.get('proficiency'),
                                 (f"{k['years']} yr{'' if k['years'] == 1 else 's'}"
                                  if isinstance(k.get('years'), (int, float)) else '')] if x)
    return ' – '.join(x for x in [k.get('name'), tail] if x)


def _style_header(ws, headers, title):
    # Row 1: title + format box
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, ncol - 1))
    t = ws.cell(row=1, column=1, value=title)
    t.font = Font(name='Arial', size=14, bold=True)
    t.alignment = _CENTER
    fb = ws.cell(row=1, column=ncol, value=FORMAT_BOX)
    fb.font = Font(name='Arial', size=9)
    fb.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 54
    # Row 2: column headers
    for i, (h, w) in enumerate(headers, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = _HDR_FONT
        c.alignment = _CENTER
        c.border = _BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'


def _add_body(ws, values):
    r = ws.max_row + 1
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = _BODY_FONT
        c.border = _BORDER
        c.alignment = Alignment(
            vertical='center',
            horizontal='center' if isinstance(v, (int, float)) else 'left',
            wrap_text=True,
        )


def build_workbook(people, subs_by_user, include_everyone):
    """people: iterable of Employee; subs_by_user: {user_id: submission_dict}."""
    wb = Workbook()
    today = date.today()

    # ── Sheet 1: Employee Skill Matrix ──
    ws1 = wb.active
    ws1.title = 'Employee Skill Matrix'
    cols1 = [
        ('Name of the Employee', 32), ('Employee ID', 12), ('Date of joining', 13),
        ('Designation', 28), ('Overall Years', 10), ('With JSAN (yrs)', 12),
        ('Years/Months', 14), ('Existing Qualifications / Competency', 40),
        ('Department', 18), ('JSAN ID', 14), ('Email', 28), ('Tools', 26),
        ('Worked company', 20), ('Role', 16), ('Customer / Project', 28),
        ('Cycle', 10), ('Submitted', 20), ('Employee note', 28),
    ]
    _style_header(ws1, cols1, 'Employee Skill Matrix')

    rows = []
    for p in people:
        sub = subs_by_user.get(p.user_id)
        if not include_everyone and not sub:
            continue
        rows.append((p, sub))

    for p, sub in rows:
        skills = (sub or {}).get('skills') or []
        tj_years, tj_label = _tenure(p.joining_date, today)
        _add_body(ws1, [
            f"{p.first_name} {p.last_name}".strip(),
            p.employee_id,
            p.joining_date.isoformat() if p.joining_date else '',
            p.designation.name if p.designation else '',
            _total_years(skills),
            tj_years,
            tj_label,
            '\n'.join(_skill_line(k) for k in skills),
            p.department.name if p.department else '',
            (p.email or '').split('@')[0],
            p.email or '',
            ', '.join(_uniq(t for k in skills for t in (k.get('tools') or []))),
            ', '.join(_uniq(k.get('company') for k in skills)),
            ', '.join(_uniq(k.get('role') for k in skills)),
            ', '.join(_uniq(k.get('customer') for k in skills)),
            (sub or {}).get('cycle_id', ''),
            (sub or {}).get('submitted_at') or '',
            (sub or {}).get('employee_note', ''),
        ])

    # ── Sheet 2: Skill details (one row per skill) ──
    ws2 = wb.create_sheet('Skill details')
    cols2 = [
        ('Name of the Employee', 32), ('Employee ID', 12), ('Designation', 28),
        ('Skill', 28), ('Category', 14), ('Level', 13), ('Years', 8),
        ('Tools', 26), ('Worked company', 20), ('Role', 16),
        ('Customer / Project', 28), ('Note', 24), ('Cycle', 10), ('Submitted', 20),
    ]
    _style_header(ws2, cols2, 'Employee Skill Matrix – Skill details')
    for p, sub in rows:
        if not sub:
            continue
        for k in (sub.get('skills') or []):
            _add_body(ws2, [
                f"{p.first_name} {p.last_name}".strip(),
                p.employee_id,
                p.designation.name if p.designation else '',
                k.get('name', ''), k.get('category', ''), k.get('proficiency', ''),
                k.get('years') if isinstance(k.get('years'), (int, float)) else None,
                ', '.join(k.get('tools') or []),
                k.get('company', ''), k.get('role', ''), k.get('customer', ''), k.get('note', ''),
                sub.get('cycle_id', ''), sub.get('submitted_at') or '',
            ])

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
