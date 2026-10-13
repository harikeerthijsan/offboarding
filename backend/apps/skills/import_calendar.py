"""Review-calendar import — ported from the source app's lib/import.ts (detectCalendar).

Accepts a CSV or XLSX with a header row containing "Cycle ID" and the period /
due / closure columns (header matching is fuzzy, as in the source). Upserts each
ReviewCycle by cycle_id. Rows with a missing/unreadable date are skipped and
reported. Dates understand the QF calendar "01-Oct-2026" style, ISO, and Excel
serials / datetime cells.
"""
import csv
import datetime
import io
import re

from .models import ReviewCycle

_MONTHS = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']


def _norm(s):
    return re.sub(r'[^a-z]', '', str(s or '').lower())


def _parse_date(v):
    if v in (None, ''):
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)):
        # Excel serial date (1900 date system).
        try:
            return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v)))
        except (ValueError, OverflowError):
            return None
    s = str(v).strip()
    if not s:
        return None
    m = re.match(r'^(\d{1,2})-([A-Za-z]{3})-(\d{4})$', s)   # 01-Oct-2026
    if m:
        mon = _MONTHS.index(m.group(2).lower()) + 1 if m.group(2).lower() in _MONTHS else 0
        if mon:
            try:
                return datetime.date(int(m.group(3)), mon, int(m.group(1)))
            except ValueError:
                return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%d %b %Y', '%d %B %Y'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.date.fromisoformat(s[:10])
    except ValueError:
        return None


def _detect_calendar(rows):
    for r in range(min(len(rows), 10)):
        cells = [_norm(c) for c in rows[r]]
        if 'cycleid' not in cells:
            continue

        def find(*keys):
            for i, c in enumerate(cells):
                if any(c.startswith(k) for k in keys):
                    return i
            return -1

        return {
            'header_row': r,
            'id': cells.index('cycleid'),
            'start': find('periodstart'),
            'end': find('periodend'),
            'employee_due': find('employeeupdatedue', 'employeedue'),
            'manager_due': find('managervalidationdue', 'managerdue'),
            'closure': find('closuretarget', 'closure'),
            'notes': find('notes'),
        }
    return None


def _read_rows(filename, content):
    """Return a list of rows (list of cell values) from a CSV or XLSX upload."""
    name = (filename or '').lower()
    if name.endswith('.csv') or name.endswith('.txt'):
        text = content.decode('utf-8-sig', errors='replace')
        return [list(r) for r in csv.reader(io.StringIO(text))]
    # Treat everything else as an Excel workbook.
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def import_calendar(filename, content):
    """Parse an upload and upsert review cycles. Returns {imported, skipped[]}."""
    rows = _read_rows(filename, content)
    cal = _detect_calendar(rows)
    if not cal:
        raise ValueError('Could not find a "Cycle ID" column. Expected columns: '
                         'Cycle ID, Period Start, Period End, Employee Update Due, '
                         'Manager Validation Due, Closure Target.')

    required = ['start', 'end', 'employee_due', 'manager_due', 'closure']
    missing_cols = [k for k in required if cal[k] < 0]
    if missing_cols:
        raise ValueError('The file is missing required date column(s): '
                         + ', '.join(missing_cols).replace('_', ' '))

    imported, skipped = 0, []
    for r in range(cal['header_row'] + 1, len(rows)):
        row = rows[r]

        def cell(idx):
            return row[idx] if 0 <= idx < len(row) else None

        cycle_id = str(cell(cal['id']) or '').strip()
        if not cycle_id:
            continue
        dates = {
            'period_start': _parse_date(cell(cal['start'])),
            'period_end': _parse_date(cell(cal['end'])),
            'employee_due': _parse_date(cell(cal['employee_due'])),
            'manager_due': _parse_date(cell(cal['manager_due'])),
            'closure_target': _parse_date(cell(cal['closure'])),
        }
        if any(v is None for v in dates.values()):
            skipped.append(f'Row {r + 1}: cycle "{cycle_id}" has a missing or unreadable date.')
            continue
        notes = str(cell(cal['notes']) or '').strip() if cal['notes'] >= 0 else ''
        ReviewCycle.objects.update_or_create(
            cycle_id=cycle_id, defaults={**dates, 'notes': notes},
        )
        imported += 1

    return {'imported': imported, 'skipped': skipped}
